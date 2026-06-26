from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


def _safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    return result.where(pd.notna(result), "")


def _autosize(worksheet) -> None:
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
        worksheet.column_dimensions[letter].width = max(10, min(max_length + 2, 42))


def _format_code_columns(worksheet) -> None:
    for header_cell in worksheet[1]:
        if header_cell.value != "代码":
            continue
        for cell in worksheet.iter_cols(
            min_col=header_cell.column,
            max_col=header_cell.column,
            min_row=1,
            max_row=worksheet.max_row,
        ):
            for code_cell in cell:
                code_cell.number_format = "000000"
                if code_cell.row == 1 or code_cell.value in (None, ""):
                    continue
                code_cell.quotePrefix = True
                code = str(code_cell.value).strip()
                code_cell.value = code.zfill(6) if code.isdigit() and len(code) < 6 else code


def _style_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.freeze_panes = "A2"
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _format_code_columns(worksheet)
    _autosize(worksheet)


def write_workbook(path: str | Path, sheets: dict[str, pd.DataFrame]) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            _safe_frame(frame).to_excel(writer, sheet_name=sheet_name, index=False)

        for worksheet in writer.book.worksheets:
            _style_sheet(worksheet)

    return output_path

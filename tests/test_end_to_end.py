from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import openpyxl
import pandas as pd

from graham_screen.cli import _market_for_frame, run_screen


ROOT = Path(__file__).resolve().parents[1]


class EndToEndTest(unittest.TestCase):
    def test_market_detection_ignores_lixinger_footer_rows(self):
        frame = pd.DataFrame({"交易所": ["hk", "数据来源于：理杏仁网站(lixinger.com)"]})

        self.assertEqual(_market_for_frame(frame), "hk")

    def test_script_entrypoint_help_loads_package(self):
        script = ROOT / "scripts" / "graham_screen.py"

        completed = subprocess.run(
            [sys.executable, str(script), "--help"],
            cwd=ROOT,
            text=True,
            capture_output=True,
            check=False,
        )

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertIn("Run A-share Graham-style", completed.stdout)

    def test_repository_csv_files_calculate_to_required_workbook(self):
        csv1 = ROOT / "公司筛选_格雷厄姆CSV1_20260622_212016.csv"
        csv2 = ROOT / "公司筛选_格雷厄姆CSV2_20260622_212026.csv"
        if not csv1.exists() or not csv2.exists():
            self.skipTest("repository CSV sample files are not present")

        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "格雷厄姆第一版筛选结果.xlsx"
            result = run_screen(csv1, csv2, output)

            self.assertTrue(output.exists())
            self.assertEqual(result.csv1_row_count, 1717)
            self.assertEqual(result.csv2_row_count, 1717)
            self.assertEqual(result.merged_company_count, 1717)
            self.assertGreaterEqual(result.non_financial_count, 1)

            workbook = openpyxl.load_workbook(output, read_only=True)
            self.assertIn("说明", workbook.sheetnames)
            self.assertIn("A档_严格通过", workbook.sheetnames)
            self.assertIn("B档_观察名单", workbook.sheetnames)
            self.assertIn("银行_候选池", workbook.sheetnames)
            self.assertIn("金融_单独观察", workbook.sheetnames)
            self.assertIn("全部评分", workbook.sheetnames)
            self.assertIn("剔除名单", workbook.sheetnames)
            self.assertIn("缺失字段", workbook.sheetnames)


if __name__ == "__main__":
    unittest.main()
